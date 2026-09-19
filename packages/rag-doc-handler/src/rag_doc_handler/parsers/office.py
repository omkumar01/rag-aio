"""Microsoft Office parsers: .docx (python-docx), .xlsx (openpyxl), .pptx (python-pptx)."""

from __future__ import annotations

import asyncio
from io import BytesIO
from typing import Any

import docx as _docx
import openpyxl as _openpyxl
from rag_core.documents import BlockKind, Document, DocumentPage, PageBlock
from rag_core.ids import new_id

from rag_doc_handler.parsers.base import BaseParser

# python-pptx ships ``py.typed`` stubs that reference ``numpy`` (3.12-only syntax);
# treat it as ``Any`` so strict mypy at py3.11 is not aborted.
_pptx: Any = __import__("pptx")

__all__ = ["DocxParser", "PptxParser", "XlsxParser"]


def _heading_level(style_name: str) -> int | None:
    """Return the heading level for a style like 'Heading 1' (1-based), else None."""
    name = style_name.lower().replace(" ", "")
    if name.startswith("heading") and len(name) > 7:
        suffix = name[7:]
        if suffix.isdigit():
            return int(suffix)
    return None


class DocxParser(BaseParser):
    """Parses ``.docx`` documents via python-docx (paragraphs + tables)."""

    @classmethod
    def supported_types(cls) -> set[str]:
        return {".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}

    async def parse(self, source: str, data: bytes) -> Document:
        return await asyncio.to_thread(self._parse_sync, source, data)

    def _parse_sync(self, source: str, data: bytes) -> Document:
        document = _docx.Document(BytesIO(data))
        blocks: list[PageBlock] = []
        texts: list[str] = []
        order = 0
        for para in document.paragraphs:
            text = para.text
            if not text:
                continue
            style = para.style.name if para.style and para.style.name else ""
            if _heading_level(style) is not None:
                kind: BlockKind = "heading"
            elif "list" in style.lower():
                kind = "list"
            else:
                kind = "text"
            blocks.append(PageBlock(id=new_id(), page_id="", kind=kind, text=text, order=order))
            order += 1
            texts.append(text)
        for table in document.tables:
            rows = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            rows = [r for r in rows if r]
            text = "\n".join(rows)
            if text:
                blocks.append(
                    PageBlock(id=new_id(), page_id="", kind="table", text=text, order=order)
                )
                order += 1
                texts.append(text)

        page = DocumentPage(id=new_id(), page_number=1, text="\n\n".join(texts), blocks=blocks)
        doc = Document(source_uri=source, text="\n\n".join(texts))
        doc.metadata.title = document.core_properties.title
        doc.metadata.author = document.core_properties.author
        doc.pages = [page]
        return self._stamp(doc, "docx")


class XlsxParser(BaseParser):
    """Parses ``.xlsx`` documents via openpyxl (one page per sheet)."""

    @classmethod
    def supported_types(cls) -> set[str]:
        return {".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

    async def parse(self, source: str, data: bytes) -> Document:
        return await asyncio.to_thread(self._parse_sync, source, data)

    def _parse_sync(self, source: str, data: bytes) -> Document:
        wb = _openpyxl.load_workbook(BytesIO(data))
        try:
            pages: list[DocumentPage] = []
            texts: list[str] = []
            for index, sheet in enumerate(wb.worksheets):
                rows = [
                    " | ".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(values_only=True)
                ]
                rows = [r for r in rows if r]
                sheet_text = "\n".join(rows)
                pages.append(
                    DocumentPage(
                        id=new_id(),
                        page_number=index + 1,
                        text=sheet_text,
                        blocks=[],
                        width=None,
                        height=None,
                    )
                )
                texts.append(sheet_text)
        finally:
            wb.close()

        doc = Document(source_uri=source, text="\n\n".join(texts))
        doc.metadata.title = wb.properties.title
        doc.pages = pages
        return self._stamp(doc, "xlsx")


class PptxParser(BaseParser):
    """Parses ``.pptx`` documents via python-pptx (one page per slide)."""

    @classmethod
    def supported_types(cls) -> set[str]:
        return {
            ".pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }

    async def parse(self, source: str, data: bytes) -> Document:
        return await asyncio.to_thread(self._parse_sync, source, data)

    def _parse_sync(self, source: str, data: bytes) -> Document:
        presentation = _pptx.Presentation(BytesIO(data))
        pages: list[DocumentPage] = []
        texts: list[str] = []
        for index, slide in enumerate(presentation.slides):
            parts = [shape.text for shape in slide.shapes if getattr(shape, "text", None)]
            try:
                notes = slide.notes_slide.notes_text_frame.text
            except Exception:  # no notes section
                notes = ""
            if notes:
                parts.append(notes)
            slide_text = "\n".join(p for p in parts if p)
            pages.append(
                DocumentPage(
                    id=new_id(),
                    page_number=index + 1,
                    text=slide_text,
                    blocks=[],
                    width=None,
                    height=None,
                )
            )
            texts.append(slide_text)

        doc = Document(source_uri=source, text="\n\n".join(texts))
        doc.metadata.title = presentation.core_properties.title
        doc.pages = pages
        return self._stamp(doc, "pptx")
