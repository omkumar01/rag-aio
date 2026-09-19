"""End-to-end parser tests using in-memory generated fixtures."""

from __future__ import annotations

import io

import docx as docx_lib
import openpyxl
import pptx
import pymupdf
from pptx.util import Inches
from rag_core.documents import PageBlock
from rag_core.ids import new_id
from rag_doc_handler.parsers.email import EmlParser
from rag_doc_handler.parsers.html import HTMLParser
from rag_doc_handler.parsers.office import DocxParser, PptxParser, XlsxParser
from rag_doc_handler.parsers.pdf import PDFParser
from rag_doc_handler.parsers.text import TextParser

HTML = """<html><head><title>My Page</title></head><body>
<script>alert('xss')</script>
<style>body{color:red}</style>
<h1>Hello</h1><h2>Sub</h2>
<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>
<img src="/img.png">
<p>Some paragraph text</p>
<a href="/a">A</a>
<a href="/b">B</a>
</body></html>"""

EML = (
    b"Subject: Hello There\r\n"
    b"From: sender@example.com\r\n"
    b"To: recipient@example.com\r\n"
    b"MIME-Version: 1.0\r\n"
    b'Content-Type: multipart/mixed; boundary="b"\r\n'
    b"\r\n"
    b"--b\r\n"
    b"Content-Type: text/plain; charset=utf-8\r\n"
    b"\r\n"
    b"This is the body.\r\n"
    b"\r\n"
    b"--b\r\n"
    b'Content-Type: application/octet-stream; name="file.txt"\r\n'
    b'Content-Disposition: attachment; filename="file.txt"\r\n'
    b"\r\n"
    b"hello attachment\r\n"
    b"--b--\r\n"
)


def _make_pdf() -> bytes:
    doc = pymupdf.open()
    p1 = doc.new_page(width=200, height=200)
    p1.insert_text((50, 50), "First page text")
    p2 = doc.new_page(width=200, height=200)
    p2.insert_text((50, 50), "Second page text")
    doc.set_metadata({"title": "Test PDF", "author": "Tester"})
    data = doc.tobytes()
    doc.close()
    return data


def _make_docx() -> bytes:
    d = docx_lib.Document()
    d.add_heading("Title Heading", level=1)
    d.add_paragraph("Body paragraph")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    d.core_properties.title = "Doc Title"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["a", "b"])
    ws1.append(["c", "d"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["x"])
    wb.properties.title = "Workbook"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_pptx() -> bytes:
    prs = pptx.Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(1)).text = "First slide content"
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(1)).text = "Second slide content"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def test_text_parser_utf8() -> None:
    doc = await TextParser().parse("note.txt", b"Hello world")
    assert doc.text == "Hello world"
    assert len(doc.pages) == 1
    assert doc.metadata.custom["encoding"] == "utf-8"
    assert doc.parser_name == "text"


async def test_text_parser_encoding_fallback() -> None:
    data = "café résumé".encode("latin-1")
    doc = await TextParser().parse("note.txt", data)
    assert "café" in doc.text
    assert doc.metadata.custom["encoding"] == "latin-1"


async def test_text_parser_markdown_metadata() -> None:
    doc = await TextParser().parse("doc.md", b"# Title\n\nbody")
    assert doc.metadata.custom["format"] == "markdown"


async def test_html_parser_extracts_title_headings_tables_images() -> None:
    doc = await HTMLParser().parse("page.html", HTML.encode("utf-8"))
    assert doc.metadata.title == "My Page"
    # Scripts and styles are stripped from the visible text.
    assert "<script>" not in doc.text
    assert "alert" not in doc.text
    kinds = [b.kind for b in doc.pages[0].blocks]
    assert "heading" in kinds
    assert "1 | 2" in doc.text
    table_blocks = [b for b in doc.pages[0].blocks if b.kind == "table"]
    assert table_blocks and "1 | 2" in table_blocks[0].text
    img_assets = [a for a in doc.assets if a.kind == "image"]
    assert img_assets and img_assets[0].content_ref == "/img.png"
    assert "Some paragraph text" in doc.text


async def test_pdf_parser_two_pages_with_metadata() -> None:
    data = _make_pdf()
    doc = await PDFParser().parse("doc.pdf", data)
    assert len(doc.pages) == 2
    assert "First page text" in doc.pages[0].text
    assert "Second page text" in doc.pages[1].text
    assert doc.metadata.title == "Test PDF"
    assert doc.metadata.author == "Tester"
    assert doc.parser_name == "pdf"


async def test_pdf_parser_includes_page_dimensions() -> None:
    data = _make_pdf()
    doc = await PDFParser().parse("doc.pdf", data)
    page = doc.pages[0]
    assert page.width == 200.0
    assert page.height == 200.0


async def test_pdf_parser_ocr_fallback_for_sparse_page() -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=100, height=100)
    page.insert_text((10, 10), "Hi")  # fewer than the 20-char default threshold
    data = doc.tobytes()
    doc.close()

    captured: dict = {}

    async def fake_ocr(document_id: str, page_number: int, image: bytes) -> list[PageBlock]:
        captured["doc_id"] = document_id
        captured["page_number"] = page_number
        captured["image_bytes"] = image
        return [
            PageBlock(id=new_id(), page_id="", kind="other", text="OCR recovered text", order=0)
        ]

    doc_result = await PDFParser(ocr_fallback=fake_ocr).parse("sparse.pdf", data)
    assert captured["page_number"] == 1
    assert len(captured["image_bytes"]) > 0
    assert any(b.text == "OCR recovered text" for b in doc_result.pages[0].blocks)


async def test_docx_parser() -> None:
    doc = await DocxParser().parse("doc.docx", _make_docx())
    assert len(doc.pages) == 1
    assert "Body paragraph" in doc.text
    assert "A | B" in doc.text
    assert doc.metadata.title == "Doc Title"
    kinds = [b.kind for b in doc.pages[0].blocks]
    assert "heading" in kinds
    assert "table" in kinds


async def test_xlsx_parser_one_page_per_sheet() -> None:
    doc = await XlsxParser().parse("data.xlsx", _make_xlsx())
    assert len(doc.pages) == 2
    assert "a | b" in doc.pages[0].text
    assert "c | d" in doc.pages[0].text
    assert "x" in doc.pages[1].text


async def test_pptx_parser_one_page_per_slide() -> None:
    doc = await PptxParser().parse("deck.pptx", _make_pptx())
    assert len(doc.pages) == 2
    assert "First slide content" in doc.pages[0].text
    assert "Second slide content" in doc.pages[1].text


async def test_eml_parser_extracts_headers_and_attachment() -> None:
    doc = await EmlParser().parse("msg.eml", EML)
    assert doc.metadata.title == "Hello There"
    assert doc.metadata.custom["from"] == "sender@example.com"
    assert doc.metadata.custom["to"] == "recipient@example.com"
    assert "This is the body." in doc.text
    assert any(a.kind == "attachment" and "file.txt" in str(a.asset_id) for a in doc.assets)
